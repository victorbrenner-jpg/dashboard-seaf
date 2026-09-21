"""Gera o Relatório MDE a partir do dataframe filtrado da tela de OB.

O arquivo parte do modelo aprovado pela área e mantém as abas de conferência:
Painel, Resumo Diário, Planilha1, RP_Consolidado e Base Fonte 500.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


COLUNAS_BASE = [
    "Número", "UG Emitente", "UG Pagadora", "Data Emissão", "Status",
    "Tipo de OB", "NE", "Credor", "Nome do Credor", "Valor", "Fonte",
    "Natureza", "Status de Envio", "RE", "PD", "GRUPO", "Elemento",
    "Despesa", "OBJETO", "DocumentoGD", "Tipo Item", "Marcador Fonte",
]


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]+", "", text.upper())


def _sheet(workbook, name: str):
    wanted = _norm(name)
    return next(sheet for sheet in workbook.worksheets if _norm(sheet.title) == wanted)


def _column(frame: pd.DataFrame, *aliases: str) -> str | None:
    index = {_norm(column): column for column in frame.columns}
    for alias in aliases:
        if _norm(alias) in index:
            return index[_norm(alias)]
    return None


def _as_value(value: object):
    return None if pd.isna(value) else value


def _money(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return pd.to_numeric(series, errors="coerce").fillna(0.0)
    text = series.fillna("").astype(str).str.replace("R$", "", regex=False).str.replace(" ", "", regex=False)
    both = text.str.contains(r"\.", regex=True) & text.str.contains(",", regex=False)
    text.loc[both] = text.loc[both].str.replace(".", "", regex=False).str.replace(",", ".", regex=False)
    only_comma = ~both & text.str.contains(",", regex=False)
    text.loc[only_comma] = text.loc[only_comma].str.replace(",", ".", regex=False)
    return pd.to_numeric(text, errors="coerce").fillna(0.0)


def _status_rp(row: pd.Series) -> str:
    text = " ".join(str(row.get(column, "")) for column in ("Status", "Tipo de OB", "Despesa", "Tipo Item"))
    text = _norm(text)
    return "Restos Não Processados (RPNP)" if "RPNP" in text or "NAOPROCESS" in text else "Restos a Pagar (RP)"


def _copy_row_style(sheet, origin: int, destination: int, last_col: int):
    for column in range(1, last_col + 1):
        source = sheet.cell(origin, column)
        target = sheet.cell(destination, column)
        target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)


def _formula_range(formula: object, last_row: int, old_last: int = 12921):
    return formula.replace(f"${old_last}", f"${last_row}") if isinstance(formula, str) and formula.startswith("=") else formula


def gerar_relatorio_mde_excel(df_ob: pd.DataFrame, modelo_path: Path) -> bytes:
    """Cria o relatório MDE completo, com fórmulas vinculadas à aba-base."""
    if df_ob is None or df_ob.empty:
        raise ValueError("Não há pagamentos para gerar o Relatório MDE.")
    if not modelo_path.exists():
        raise FileNotFoundError("Modelo MDE não encontrado.")

    source = df_ob.copy()
    col_data = _column(source, "Data Emissão", "Data Emissao", "Data")
    col_valor = _column(source, "Valor")
    if not col_data or not col_valor:
        raise ValueError("A base precisa conter as colunas Data Emissão e Valor.")

    dates = pd.to_datetime(source[col_data], errors="coerce", dayfirst=True)
    amounts = _money(source["Valor_Limpo"] if "Valor_Limpo" in source else source[col_valor])
    expense = source["Despesa_Tratada"] if "Despesa_Tratada" in source else source.get("Despesa", "CORRENTE")
    expense = expense.fillna("CORRENTE").astype(str).str.upper()
    marker_col = _column(source, "Marcador_Fonte_Tratado", "Marcador Fonte", "Marcador de Fonte")

    # A aba-base usa os mesmos 22 campos do modelo. As colunas tratadas ficam
    # apenas no dataframe; no Excel, Despesa recebe a classificação que alimenta
    # todas as SOMASES do painel.
    base_frame = pd.DataFrame(index=source.index)
    for header in COLUNAS_BASE:
        origin = _column(source, header)
        base_frame[header] = source[origin] if origin else None
    base_frame["Data Emissão"] = dates
    base_frame["Valor"] = amounts
    base_frame["Despesa"] = expense
    if marker_col:
        base_frame["Marcador Fonte"] = source[marker_col].fillna("NÃO INFORMADO").astype(str)
    base_frame = base_frame.where(pd.notna(base_frame), None)

    workbook = load_workbook(modelo_path)
    painel = _sheet(workbook, "Painel")
    diario = _sheet(workbook, "Resumo Diario")
    planilha_rp = _sheet(workbook, "Planilha1")
    rp_sheet = _sheet(workbook, "RP_Consolidado")
    base_sheet = _sheet(workbook, "Base Fonte 500")

    # Mantém a aba de origem editável: ao alterar Valor/Data/Despesa/Marcador,
    # o Painel recalcula no Excel pelas fórmulas preservadas no modelo.
    if base_sheet.max_row > 1:
        base_sheet.delete_rows(2, base_sheet.max_row - 1)
    for col, header in enumerate(COLUNAS_BASE, start=1):
        base_sheet.cell(1, col).value = header
    for values in base_frame.itertuples(index=False, name=None):
        base_sheet.append(list(values))
    last_base = max(2, len(base_frame) + 1)
    base_sheet.auto_filter.ref = f"A1:V{last_base}"
    base_sheet.freeze_panes = "A2"

    # Todas as fórmulas do Painel passam a apontar para a extensão real da base.
    for row in painel.iter_rows():
        for cell in row:
            if cell.__class__.__name__ != "MergedCell":
                cell.value = _formula_range(cell.value, last_base)

    # Resumo diário é refeito para refletir exatamente os filtros ativos, sem
    # perder a vinculação com a aba Base Fonte 500 do arquivo exportado.
    dates_valid = dates.dropna().dt.normalize()
    unique_dates = sorted(dates_valid.unique())
    if diario.max_row > 3:
        for row in diario.iter_rows(min_row=4, max_row=diario.max_row, min_col=1, max_col=6):
            for cell in row:
                cell.value = None
    daily_start = 4
    for offset, date_value in enumerate(unique_dates):
        row = daily_start + offset
        _copy_row_style(diario, 4, row, 6)
        diario.cell(row, 1).value = pd.Timestamp(date_value).to_pydatetime()
        diario.cell(row, 2).value = f'=COUNTIFS(\'Base Fonte 500\'!$D$2:$D${last_base},A{row})'
        diario.cell(row, 3).value = f'=SUMIFS(\'Base Fonte 500\'!$J$2:$J${last_base},\'Base Fonte 500\'!$D$2:$D${last_base},A{row})'
        for col, kind in ((4, "CORRENTE"), (5, "RP"), (6, "DEA")):
            diario.cell(row, col).value = f'=SUMIFS(\'Base Fonte 500\'!$J$2:$J${last_base},\'Base Fonte 500\'!$D$2:$D${last_base},A{row},\'Base Fonte 500\'!$R$2:$R${last_base},"{kind}")'
    total_row = max(104, daily_start + len(unique_dates))
    _copy_row_style(diario, 104, total_row, 6)
    diario.cell(total_row, 1).value = "TOTAL GERAL"
    for col in range(2, 7):
        letter = chr(64 + col)
        diario.cell(total_row, col).value = f"=SUM({letter}{daily_start}:{letter}{total_row - 1})"
    diario.auto_filter.ref = f"A3:F{max(3, total_row - 1)}"

    # Detalhamento de RP e a planilha de conferência por MDE 1001.
    rp_mask = expense.eq("RP")
    rp_data = source.loc[rp_mask].copy()
    if rp_sheet.max_row > 1:
        rp_sheet.delete_rows(2, rp_sheet.max_row - 1)
    rp_headers = [cell.value for cell in rp_sheet[1]]
    for col, header in enumerate(rp_headers, start=1):
        rp_sheet.cell(1, col).value = header
    month_pt = {1: "jan", 2: "fev", 3: "mar", 4: "abr", 5: "mai", 6: "jun", 7: "jul", 8: "ago", 9: "set", 10: "out", 11: "nov", 12: "dez"}
    for idx, (_, row) in enumerate(rp_data.iterrows(), start=2):
        date_value = pd.to_datetime(row.get(col_data), errors="coerce", dayfirst=True)
        values = [
            month_pt.get(date_value.month, "") if pd.notna(date_value) else "",
            row.get(_column(source, "Fonte"), None), amounts.loc[row.name],
            row.get(marker_col, "NÃO INFORMADO") if marker_col else "NÃO INFORMADO",
            row.get(_column(source, "GRUPO"), None), row.get(_column(source, "Elemento"), None),
            row.get(_column(source, "Tipo Item"), None), row.get(_column(source, "Credor", "Nome do Credor"), None),
            row.get(_column(source, "OBJETO"), None), row.get(_column(source, "NE"), None),
            row.get(_column(source, "PD"), None), row.get(_column(source, "Número"), None),
            row.get(_column(source, "DocumentoGD"), None), _status_rp(row),
            date_value.to_pydatetime() if pd.notna(date_value) else None,
        ]
        for col, value in enumerate(values, start=1):
            rp_sheet.cell(idx, col).value = _as_value(value)
    last_rp = max(2, len(rp_data) + 1)
    rp_sheet.auto_filter.ref = f"A1:O{last_rp}"
    for row in planilha_rp.iter_rows():
        for cell in row:
            if cell.__class__.__name__ != "MergedCell":
                cell.value = _formula_range(cell.value, last_rp, old_last=783)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
